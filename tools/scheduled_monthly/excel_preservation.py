from __future__ import annotations

"""Keep exported workbook formulas/content/styles while normalizing date display."""

import os
import shutil
import subprocess
import tempfile
from copy import copy
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


DATE_COLUMNS = ("C", "D", "H")
DATE_NUMBER_FORMAT = "yyyy/m/d"


def _soffice_path() -> str:
    configured = os.getenv("SOFFICE_PATH", "").strip()
    candidates = [
        configured,
        shutil.which("soffice") or "",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    raise RuntimeError("找不到 LibreOffice soffice，無法在保留公式與格式的前提下轉換 XLS")


def _set_date_formats(path: str) -> None:
    workbook = load_workbook(path, data_only=False)
    for worksheet in workbook.worksheets:
        for column in DATE_COLUMNS:
            for cell in worksheet[column][1:]:
                if cell.value not in (None, ""):
                    cell.number_format = DATE_NUMBER_FORMAT
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)


def save_preserved_excel(content: bytes, output_path: str) -> str:
    """Convert XLS/XLSX without flattening formulas, then only format C/D/H as dates."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir)
        if content[:2] == b"PK":
            converted = temp / "source.xlsx"
            converted.write_bytes(content)
        elif content[:4] == b"\xd0\xcf\x11\xe0":
            source = temp / "source.xls"
            converted_dir = temp / "converted"
            profile_dir = temp / "libreoffice-profile"
            source.write_bytes(content)
            converted_dir.mkdir()
            command = [
                _soffice_path(),
                f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
                "--headless", "--nologo", "--nodefault", "--nolockcheck",
                "--nofirststartwizard", "--convert-to", "xlsx",
                "--outdir", str(converted_dir), str(source),
            ]
            result = subprocess.run(command, check=False, capture_output=True, text=True, timeout=120)
            converted = converted_dir / "source.xlsx"
            if result.returncode != 0 or not converted.exists():
                detail = (result.stderr or result.stdout or "未知錯誤").strip()
                raise RuntimeError(f"LibreOffice 轉換 XLS 失敗：{detail}")
        else:
            raise RuntimeError("下載內容不是可辨識的 XLS/XLSX")
        shutil.copy2(converted, output)
    _set_date_formats(str(output))
    return str(output)


def _copy_cell(source, target) -> None:
    value = source.value
    if source.data_type == "f" and isinstance(value, str):
        try:
            value = Translator(value, origin=source.coordinate).translate_formula(target.coordinate)
        except Exception:
            pass
    target.value = value
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def _row_key(values: Iterable[object]) -> tuple[object, ...]:
    normalized = []
    for value in values:
        if hasattr(value, "isoformat"):
            value = value.isoformat()
        normalized.append(value)
    return tuple(normalized)


def filter_workbook_rows(source_path: str, output_path: str, order_numbers: set[str]) -> str:
    """Keep the header and selected order rows while retaining cell formulas and styles."""
    source_book = load_workbook(source_path, data_only=False)
    source_sheet = source_book.active
    selected_rows = [
        row for row in range(2, source_sheet.max_row + 1)
        if str(source_sheet.cell(row, 2).value or "").strip() in order_numbers
    ]
    source_snapshot = load_workbook(source_path, data_only=False).active
    if source_sheet.max_row > 1:
        source_sheet.delete_rows(2, source_sheet.max_row - 1)
    for target_row, source_row in enumerate(selected_rows, start=2):
        source_sheet.row_dimensions[target_row].height = source_snapshot.row_dimensions[source_row].height
        for column in range(1, source_snapshot.max_column + 1):
            _copy_cell(source_snapshot.cell(source_row, column), source_sheet.cell(target_row, column))
    source_book.save(output_path)
    _set_date_formats(output_path)
    return output_path


def merge_workbook_files(source_paths: list[str], output_path: str) -> str:
    """Append data rows from matching workbooks, preserving formulas/styles and removing duplicates."""
    if not source_paths:
        raise RuntimeError("沒有可合併的 Excel 來源檔")
    if Path(source_paths[0]).resolve() != Path(output_path).resolve():
        shutil.copy2(source_paths[0], output_path)
    target_book = load_workbook(output_path, data_only=False)
    target_sheet = target_book.active
    target_values = load_workbook(output_path, data_only=True).active
    seen = {
        _row_key(target_values.cell(row, column).value for column in range(1, target_values.max_column + 1))
        for row in range(2, target_values.max_row + 1)
    }
    for source_path in source_paths[1:]:
        source_book = load_workbook(source_path, data_only=False)
        source_sheet = source_book.active
        value_sheet = load_workbook(source_path, data_only=True).active
        for source_row in range(2, source_sheet.max_row + 1):
            key = _row_key(value_sheet.cell(source_row, column).value for column in range(1, value_sheet.max_column + 1))
            if key in seen:
                continue
            seen.add(key)
            target_row = target_sheet.max_row + 1
            target_sheet.row_dimensions[target_row].height = source_sheet.row_dimensions[source_row].height
            for column in range(1, source_sheet.max_column + 1):
                _copy_cell(source_sheet.cell(source_row, column), target_sheet.cell(target_row, column))
    target_book.save(output_path)
    _set_date_formats(output_path)
    return output_path
