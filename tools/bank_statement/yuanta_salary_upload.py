"""Find Yuanta salary workbooks in Drive and upload them with the local Agent Chrome."""

from __future__ import annotations

from contextlib import ExitStack

import argparse
import io
import json
import re
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from services.google_auth import get_drive_service
from tools.bank_statement.accounts import DEFAULT_ACCOUNTS_FILE, load_account
from tools.bank_statement.fubon_agent import ensure_login, run_download, report_query_start
from tools.bank_statement.fubon_agent import logout_quietly as fubon_logout_quietly
from tools.bank_statement.fubon_transfer_common import (
    choose_immediate_date,
    choose_saved_destination,
    choose_source_account,
    fill_row_inputs,
    open_transfer_form,
    transfer_page_stage,
    wait_user_completed_transfer,
)
from tools.bank_statement.native_yuanta import (
    YUANTA_HOST,
    YUANTA_URL,
    bind_yuanta_page,
    logout_yuanta_quietly,
    wait_for_yuanta_login,
)
from tools.invoice_center.chrome_cdp import connect_existing_chrome
from tools.invoice_center.invoice_archive import AREA_FOLDER_NAMES, MONTHLY_ROOT_FOLDER_ID
from tools.staff_payroll import ROOT_FOLDER_ID as INTERNAL_SALARY_ROOT_ID


FOLDER_MIME = "application/vnd.google-apps.folder"
SUPPORTED_AREAS = tuple(AREA_FOLDER_NAMES)


@dataclass(frozen=True)
class SalaryFile:
    file_id: str
    name: str
    amount: int
    kind: str
    summary: str


def _escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("'", "\\'")


def _children(drive: Any, parent_id: str, *, folders_only: bool = False) -> list[dict[str, Any]]:
    query = f"'{_escape(parent_id)}' in parents and trashed=false"
    if folders_only:
        query += f" and mimeType='{FOLDER_MIME}'"
    result = drive.files().list(
        q=query,
        fields="files(id,name,mimeType,modifiedTime)",
        pageSize=1000,
        supportsAllDrives=True,
        includeItemsFromAllDrives=True,
    ).execute()
    return list(result.get("files", []))


def _find_exact(drive: Any, parent_id: str, name: str, *, folder: bool = False) -> dict[str, Any] | None:
    rows = [item for item in _children(drive, parent_id, folders_only=folder) if item.get("name") == name]
    rows.sort(key=lambda item: str(item.get("modifiedTime") or ""), reverse=True)
    return rows[0] if rows else None


def _find_area_folder(drive: Any, parent_id: str, area: str) -> dict[str, Any]:
    expected = AREA_FOLDER_NAMES.get(area, f"{area}專員")
    folders = _children(drive, parent_id, folders_only=True)
    exact = [item for item in folders if item.get("name") == expected]
    if exact:
        return exact[0]
    matches = [
        item for item in folders
        if re.sub(r"^[\d.\s]+", "", str(item.get("name") or "")) == f"{area}專員"
    ]
    if len(matches) != 1:
        raise FileNotFoundError(f"找不到唯一的地區資料夾：{expected}")
    return matches[0]


def _number(value: object) -> Decimal:
    if value is None or value == "" or isinstance(value, bool):
        return Decimal(0)
    try:
        return Decimal(str(value).strip().replace(",", "").replace("$", ""))
    except InvalidOperation:
        return Decimal(0)


def xlsx_total(data: bytes) -> int:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    worksheet = workbook["元大"] if "元大" in workbook.sheetnames else workbook.active
    total = sum(
        (_number(row[0]) for row in worksheet.iter_rows(min_row=3, min_col=4, max_col=4, values_only=True)),
        Decimal(0),
    )
    if total != total.to_integral_value():
        raise ValueError(f"薪資檔金額不是整數：{total}")
    return int(total)


def _download(drive: Any, file_id: str) -> bytes:
    return drive.files().get_media(fileId=file_id, supportsAllDrives=True).execute()


def _salary_file(drive: Any, item: dict[str, Any], kind: str, summary: str) -> SalaryFile:
    return SalaryFile(
        file_id=str(item["id"]),
        name=str(item["name"]),
        amount=xlsx_total(_download(drive, str(item["id"]))),
        kind=kind,
        summary=summary,
    )


def find_salary_files(
    drive: Any,
    period: str,
    area: str,
    *,
    include_contract: bool = True,
    include_internal: bool = False,
) -> list[SalaryFile]:
    """Resolve the exact Drive paths requested by the finance workflow."""
    if not re.fullmatch(r"\d{6}-[12]", period):
        raise ValueError("期別格式須為 YYYYMM-1 或 YYYYMM-2")
    if area not in SUPPORTED_AREAS:
        raise ValueError(f"不支援的地區：{area}")
    if not include_contract and not include_internal:
        raise ValueError("請至少選擇專員承攬或內勤薪資")

    files: list[SalaryFile] = []
    year, yyyymm = period[:4], period[:6]

    if include_contract:
        year_folder = _find_exact(drive, MONTHLY_ROOT_FOLDER_ID, f"{year}專員承攬服務費", folder=True)
        if not year_folder:
            raise FileNotFoundError(f"找不到資料夾：{year}專員承攬服務費")
        area_folder = _find_area_folder(drive, str(year_folder["id"]), area)
        period_folder = _find_exact(drive, str(area_folder["id"]), period, folder=True)
        if not period_folder:
            raise FileNotFoundError(f"找不到期別資料夾：{period}")

        fee_name = f"{period}元大承攬費-{area}.xlsx"
        deposit_name = f"{period}元大工具包押金-{area}.xlsx"
        fee = _find_exact(drive, str(period_folder["id"]), fee_name)
        deposit = _find_exact(drive, str(period_folder["id"]), deposit_name)
        if not fee:
            raise FileNotFoundError(f"找不到必要檔案：{fee_name}")
        files.append(_salary_file(drive, fee, "專員承攬費", "員工酬勞"))
        if deposit:
            files.append(_salary_file(drive, deposit, "工具包押金", "費用"))

    if include_internal:
        year_folder = _find_exact(drive, INTERNAL_SALARY_ROOT_ID, f"{year}薪資", folder=True)
        if not year_folder:
            raise FileNotFoundError(f"找不到內勤資料夾：{year}薪資")
        area_folder = _find_exact(drive, str(year_folder["id"]), area, folder=True)
        if not area_folder:
            raise FileNotFoundError(f"找不到內勤地區資料夾：{area}")
        candidates = (f"{yyyymm}元大帳戶-{area}.xlsx", f"{year}元大帳戶-{area}.xlsx")
        internal = None
        for name in candidates:
            internal = _find_exact(drive, str(area_folder["id"]), name)
            if internal:
                break
        if not internal:
            raise FileNotFoundError(f"找不到內勤元大檔：{' 或 '.join(candidates)}")
        files.append(_salary_file(drive, internal, "內勤薪資", "薪資"))

    return files


def _all_frames(page: Any) -> list[Any]:
    # Agent Chrome 可能同時開著富邦、Google 等分頁；元大操作只能掃描元大頁面，
    # 否則可能在其他分頁的下拉選單上等待到逾時。
    pages = [
        candidate
        for candidate in page.context.pages
        if not candidate.is_closed() and YUANTA_HOST in candidate.url
    ]
    if not pages and not page.is_closed():
        pages = [page]
    return [frame for candidate in pages for frame in candidate.frames]


def _visible(locator: Any) -> Any | None:
    try:
        for index in range(locator.count()):
            item = locator.nth(index)
            if item.is_visible():
                return item
    except Exception:
        return None
    return None


def _normalized_text(value: object) -> str:
    """Normalize bank labels whose spacing changes across browsers/OSes."""
    return re.sub(r"\s+", "", str(value or ""))


def _click_text(page: Any, text: str, timeout_seconds: int = 30) -> None:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        for frame in _all_frames(page):
            for exact in (True, False):
                item = _visible(frame.get_by_text(text, exact=exact))
                if item:
                    item.click()
                    return
        page.wait_for_timeout(500)
    raise RuntimeError(f"元大頁面找不到「{text}」")


def _select_option_containing(
    page: Any, option_text: str, label_hint: str = "", timeout_seconds: int = 30
) -> None:
    expected = _normalized_text(option_text)
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        for frame in _all_frames(page):
            selects = frame.locator("select")
            for index in range(selects.count()):
                select = selects.nth(index)
                try:
                    if not select.is_visible():
                        continue
                    if label_hint:
                        surrounding = select.locator("xpath=ancestor::*[self::tr or self::div][1]")
                        if surrounding.count() and label_hint not in surrounding.inner_text(timeout=1_000):
                            continue
                    # 一次從 DOM 取得全部 option，避免逐項 inner_text 各自等待逾時。
                    options = select.evaluate(
                        "el => Array.from(el.options).map(o => ({text: o.textContent || '', value: o.value}))"
                    )
                except Exception:
                    continue
                for option in options:
                    if expected not in _normalized_text(option.get("text")):
                        continue
                    try:
                        select.select_option(value=str(option.get("value", "")), timeout=3_000)
                        return
                    except Exception:
                        break
        page.wait_for_timeout(250)
    raise RuntimeError(f"元大頁面找不到選項「{option_text}」")


def _set_file(page: Any, path: Path) -> None:
    for frame in _all_frames(page):
        item = _visible(frame.locator("input[type=file]"))
        if item:
            item.set_input_files(str(path))
            return
    raise RuntimeError("元大頁面找不到檔案上傳欄位")


def upload_salary_files(files: list[SalaryFile], area: str, cdp_url: str, accounts_file: Path) -> None:
    account = load_account("yuanta", area, accounts_file.expanduser())
    drive = get_drive_service()
    with tempfile.TemporaryDirectory(prefix="yuanta-salary-") as tmp:
        staged: list[tuple[SalaryFile, Path]] = []
        for item in files:
            path = Path(tmp) / item.name
            path.write_bytes(_download(drive, item.file_id))
            staged.append((item, path))

        with sync_playwright() as playwright:
            _browser, context = connect_existing_chrome(playwright, cdp_url)
            pages = [page for page in context.pages if not page.is_closed() and YUANTA_HOST in page.url]
            page = pages[-1] if pages else context.new_page()
            if YUANTA_HOST not in page.url:
                page.goto(YUANTA_URL, wait_until="domcontentloaded", timeout=60_000)
            bind_yuanta_page(page)
            wait_for_yuanta_login(account, timeout=300)
            try:
                _upload_staged(page, account, staged)
            finally:
                logout_yuanta_quietly()  # 送審完（或失敗）一律登出


def _upload_staged(page: Any, account: Any, staged: list) -> None:
    """逐檔上傳、選付款帳號（末五碼）、送審。"""
    for item, path in staged:
        print(f"上傳中：{item.name}（NT$ {item.amount:,}）", flush=True)
        _click_text(page, "收付款服務")
        _click_text(page, "薪資付款")
        _click_text(page, "薪資臺幣整批付款")
        _select_option_containing(page, "臺幣薪轉格式Excel")
        _set_file(page, path)
        _click_text(page, "確定")
        # 帳號顯示名稱可能是「元大大同-綜合活期-...」等不同格式；
        # 以各地區設定的帳號末五碼選取，避免依賴銀行頁面的標點文字。
        payment_account = account.bank_account[-5:] if account.bank_account else "元大大同"
        _select_option_containing(page, payment_account)
        _select_option_containing(page, item.summary, "摘要")
        _click_text(page, "確定匯入")
        _click_text(page, "送審")
        _click_text(page, "確認送審")
        print(f"已送待放行：{item.name}", flush=True)


def transfer_total_from_fubon(
    amount: int,
    area: str,
    period: str,
    cdp_url: str,
    accounts_file: Path,
) -> None:
    """Fill a Fubon transfer to the saved Yuanta account, then sync today's statement."""
    if amount <= 0:
        raise ValueError("富邦匯往元大的總金額必須大於 0")
    account = load_account("fubon", area, accounts_file.expanduser())
    destination = f"元大{area}"
    with sync_playwright() as playwright, ExitStack() as cleanup:
        _browser, context = connect_existing_chrome(playwright, cdp_url)
        page = ensure_login(context, account)
        cleanup.callback(fubon_logout_quietly, context, page)  # 轉帳完（或失敗）一律登出富邦
        stage = transfer_page_stage(page, str(amount))
        if stage in {"confirmation", "completed"}:
            print(f"偵測到富邦目前停在{stage}階段，從現有畫面接續。", flush=True)
        else:
            page = open_transfer_form(page)
            choose_source_account(page, area, account.bank_account)
            choose_saved_destination(page, destination)
            page.wait_for_timeout(800)
            fill_row_inputs(page, "轉帳金額", [str(amount)])
            choose_immediate_date(page)
            memo = f"{period}薪資"
            fill_row_inputs(page, "給自己", [memo])
            fill_row_inputs(page, "給對方", [memo])
            print(f"富邦已填入：NT$ {amount:,}，轉入常用帳號「{destination}」。", flush=True)
            print("請人工核對並完成富邦交易；確認成功頁後系統才會繼續。", flush=True)
        wait_user_completed_transfer(
            page,
            require_completed_at=True,
            expected_amount=str(amount),
        )
        try:
            today = datetime.now().date()
            print(f"富邦轉帳完成，開始同步 {area} 富邦明細（財報最後一列日期～今天）。", flush=True)
            run_download(context, page, account, report_query_start(account, today), today)
            print(f"{area} 富邦明細更新完成。", flush=True)
        except Exception as exc:
            raise RuntimeError(
                "富邦轉帳已完成，但明細更新失敗；請勿重複匯款，請依日誌補做明細更新："
                f"{exc}"
            ) from exc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="元大薪資檔上傳待放行")
    parser.add_argument("--area", required=True, choices=SUPPORTED_AREAS)
    parser.add_argument("--period", required=True, help="YYYYMM-1 或 YYYYMM-2")
    parser.add_argument("--no-contract", action="store_true", help="不找專員承攬費／工具包押金")
    parser.add_argument("--internal", action="store_true", help="一併上傳內勤薪資元大帳戶檔")
    parser.add_argument("--fubon-transfer", action="store_true", help="先由富邦匯總額到該地區元大帳戶")
    parser.add_argument("--expected-plan", help="畫面預覽確認過的檔名與金額 JSON")
    parser.add_argument("--cdp-url", default="http://127.0.0.1:9222")
    parser.add_argument("--accounts-file", type=Path, default=DEFAULT_ACCOUNTS_FILE)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    drive = get_drive_service()
    files = find_salary_files(
        drive,
        args.period,
        args.area,
        include_contract=not args.no_contract,
        include_internal=args.internal,
    )
    if args.expected_plan:
        expected = json.loads(args.expected_plan)
        actual = [{"name": item.name, "amount": item.amount} for item in files]
        if expected != actual:
            raise RuntimeError(
                "Drive 檔案或金額已和畫面預覽不同，為安全起見停止執行；"
                "請回到 Tool System 重新檢查本次檔案與金額"
            )
    for item in files:
        print(f"找到：{item.kind}｜{item.name}｜NT$ {item.amount:,}", flush=True)
    total = sum(item.amount for item in files)
    print(f"合計：NT$ {total:,}", flush=True)
    transferred = False
    if args.fubon_transfer:
        transfer_total_from_fubon(total, args.area, args.period, args.cdp_url, args.accounts_file)
        transferred = True
    try:
        upload_salary_files(files, args.area, args.cdp_url, args.accounts_file)
    except Exception as exc:
        if transferred:
            raise RuntimeError(
                "富邦轉帳及明細更新已完成，但元大薪資檔上傳失敗；請勿重複執行富邦轉帳，"
                f"請依日誌續做元大上傳：{exc}"
            ) from exc
        raise
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
