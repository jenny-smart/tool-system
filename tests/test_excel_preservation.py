from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from tools.scheduled_monthly.excel_preservation import (
    DATE_NUMBER_FORMAT,
    filter_workbook_rows,
    merge_workbook_files,
    save_preserved_excel,
)


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["狀態", "訂單編號", "購買日期", "付款日期", "E", "F", "G", "服務日期"] + [f"欄{i}" for i in range(9, 38)] + ["加收金流"])
    for order_no, row_number in rows:
        sheet.append(["已加收", order_no, "2026-10-01", f"=AK{row_number}", "保留內容", "", "", "2026-10-03"] + [""] * 28 + ["ATM"])
    sheet["E2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    sheet.row_dimensions[2].height = 24
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_save_preserves_formula_content_and_style_and_only_formats_dates(tmp_path):
    output = tmp_path / "orders.xlsx"
    save_preserved_excel(workbook_bytes([("LC1-1", 2)]), str(output))
    sheet = load_workbook(output, data_only=False).active
    assert sheet["D2"].value == "=AK2"
    assert sheet["E2"].value == "保留內容"
    assert sheet["E2"].fill.fgColor.rgb == "00FFFF00"
    assert sheet.row_dimensions[2].height == 24
    assert [sheet[f"{column}2"].number_format for column in ("C", "D", "H")] == [DATE_NUMBER_FORMAT] * 3
    assert sheet["E2"].number_format == "General"


def test_filter_and_merge_keep_formulas_and_translate_row_references(tmp_path):
    source = tmp_path / "source.xlsx"
    save_preserved_excel(workbook_bytes([("LC1-1", 2), ("LC2-1", 3)]), str(source))
    filtered = tmp_path / "filtered.xlsx"
    filter_workbook_rows(str(source), str(filtered), {"LC2-1"})
    filtered_sheet = load_workbook(filtered, data_only=False).active
    assert filtered_sheet["B2"].value == "LC2-1"
    assert filtered_sheet["D2"].value == "=AK2"

    other = tmp_path / "other.xlsx"
    save_preserved_excel(workbook_bytes([("LC3-1", 2)]), str(other))
    merged = tmp_path / "merged.xlsx"
    merge_workbook_files([str(filtered), str(other)], str(merged))
    merged_sheet = load_workbook(merged, data_only=False).active
    assert [merged_sheet[f"B{row}"].value for row in (2, 3)] == ["LC2-1", "LC3-1"]
    assert [merged_sheet[f"D{row}"].value for row in (2, 3)] == ["=AK2", "=AK3"]
